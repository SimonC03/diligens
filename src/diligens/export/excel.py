"""The DD pack as an Excel workbook. Swedish labels, TSEK, live formulas for totals.

Values are written in TSEK (SEK / 1000, unrounded), so formulas compute exactly and the
number format only controls display. Signs follow DOMAIN.md section 3: revenue positive and
costs negative in the income statement; assets, equity and liabilities positive in the
balance sheet. Every sheet has a "Källa" column.
"""

import datetime as dt
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from diligens.analysis.kpis import nwc_lines
from diligens.analysis.mapping import BAS_LINES
from diligens.analysis.statements import COMPUTED_RESULT
from diligens.models import REVIEW_CONFIDENCE, AnalysisResult, BridgeItem, FlagCategory

SHEETS = [
    "Sammanfattning",
    "RR månad",
    "BR månad",
    "EBITDA brygga",
    "Rörelsekapital",
    "Nettoskuld",
    "Flaggade poster",
    "DD frågor",
    "Avstämning",
    "Källor",
]

TSEK = "#,##0.0"
SEK = "#,##0.00"
PERCENT = "0.0%"
CONFIDENCE = "0.00"
THOUSAND = Decimal(1000)

CATEGORY_LABELS: dict[FlagCategory, str] = {
    "engangspost": "Engångspost",
    "agarrelaterad": "Ägarrelaterad",
    "narstaende": "Närstående",
    "periodisering": "Periodisering",
    "avvikelse": "Avvikelse",
    "normal": "Normal",
}
NWC_LABELS = {
    "Leverantörsskulder och kortfristiga skulder till kreditinstitut": "Leverantörsskulder",
}
STATUS_ORDER = ["Justerad", "Kräver granskning", "Ej justerad, fråga", "Utanför perioden", "Normal"]

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDE3EA")
_TOTAL_FONT = Font(bold=True)

# Income statement layout: BAS lines and subtotals (label, parts) in presentation order.
RR_LAYOUT: list[tuple[str, list[str] | None]] = [
    ("Nettoomsättning", None),
    ("Råvaror, handelsvaror och underentreprenader", None),
    ("Bruttoresultat", ["Nettoomsättning", "Råvaror, handelsvaror och underentreprenader"]),
    ("Övriga rörelseintäkter", None),
    ("Övriga externa kostnader", None),
    ("Personalkostnader", None),
    ("Övriga rörelsekostnader", None),
    (
        "EBITDA",
        [
            "Bruttoresultat",
            "Övriga rörelseintäkter",
            "Övriga externa kostnader",
            "Personalkostnader",
            "Övriga rörelsekostnader",
        ],
    ),
    ("Av och nedskrivningar", None),
    ("EBIT", ["EBITDA", "Av och nedskrivningar"]),
    ("Finansiella poster", None),
    ("Bokslutsdispositioner", None),
    ("Resultat före skatt", ["EBIT", "Finansiella poster", "Bokslutsdispositioner"]),
    ("Skatt", None),
    ("Årets resultat", ["Resultat före skatt", "Skatt"]),
]


def _tsek(value: Decimal | None) -> Decimal | None:
    return None if value is None else value / THOUSAND


def _col(index: int) -> str:
    return get_column_letter(index)


def _header(ws: Worksheet, labels: list[str]) -> None:
    ws.append(labels)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _bold(ws: Worksheet, row: int) -> None:
    for cell in ws[row]:
        cell.font = _TOTAL_FONT


def _format(ws: Worksheet, columns: range | list[int], rows: range, number_format: str) -> None:
    for row in rows:
        for column in columns:
            ws.cell(row=row, column=column).number_format = number_format


def _widths(ws: Worksheet, fixed: dict[int, int] | None = None) -> None:
    fixed = fixed or {}
    for index, column in enumerate(ws.iter_cols(), start=1):
        if index in fixed:
            width = fixed[index]
        else:
            lengths = [
                len(str(c.value))
                for c in column
                if c.value is not None and not str(c.value).startswith("=")
            ]
            width = min(max([12, *lengths]) + 2, 60)
        ws.column_dimensions[_col(index)].width = width


def _income_statement(ws: Worksheet, result: AnalysisResult) -> tuple[dict[str, int], str]:
    """Write RR månad; return label to row and the period (LTM or YTD) column letter."""
    months = result.statements.months
    bridge = result.bridge
    period_title = (
        f"{'LTM' if bridge.basis == 'LTM' else 'Hittills i år'} "
        f"{bridge.period[0]} till {bridge.period[1]}"
    )
    _header(ws, ["Rad", *months, period_title, "Källa"])
    first = months.index(bridge.period[0]) + 2
    last = months.index(bridge.period[1]) + 2
    period_col = len(months) + 2
    lines = {line.line: line for line in result.statements.income}
    # Tax only (8900-8989): the closing entry on 8990-8999 moves the result to equity and
    # would read as a tax charge.
    tax_accounts = [a for a in result.statements.account_movements if 8900 <= int(a) <= 8989]
    tax = {
        m: sum((result.statements.account_movements[a][m] for a in tax_accounts), Decimal(0))
        for m in months
    }
    sources = {label: line.source_ref for label, line in lines.items()}
    sources["Skatt"] = (
        "SIE #TRANS per verifikationsdatum, konton 8900–8989. Bokslutets årets resultat "
        "(8990–8999) visas inte."
    )

    rows: dict[str, int] = {}
    for label, parts in RR_LAYOUT:
        row = ws.max_row + 1
        rows[label] = row
        ws.cell(row=row, column=1, value=label)
        values = tax if label == "Skatt" else (lines[label].values if parts is None else {})
        for i, month in enumerate(months, start=2):
            letter = _col(i)
            if parts is None:
                ws.cell(row=row, column=i, value=_tsek(-values[month]))
            else:
                ws.cell(
                    row=row, column=i, value="=" + "+".join(f"{letter}{rows[p]}" for p in parts)
                )
        ws.cell(
            row=row,
            column=period_col,
            value=f"=SUM({_col(first)}{row}:{_col(last)}{row})",
        )
        source = sources[label] if parts is None else "Formel: " + " + ".join(parts)
        ws.cell(row=row, column=period_col + 1, value=source)
        if parts is not None:
            _bold(ws, row)
    _format(ws, range(2, period_col + 1), range(2, ws.max_row + 1), TSEK)
    ws.freeze_panes = "B2"
    _widths(ws, {1: 44})
    return rows, _col(period_col)


def _balance_sheet(ws: Worksheet, result: AnalysisResult) -> None:
    months = result.statements.months
    _header(ws, ["Rad", *months, "Källa"])
    source_col = len(months) + 2
    asset_lines = {
        line for first, _, line, statement in BAS_LINES if statement == "BR" and first < 2000
    }
    lines = result.statements.balance

    def write(label: str, values: dict[str, Decimal], sign: int, source: str) -> int:
        row = ws.max_row + 1
        ws.cell(row=row, column=1, value=label)
        for i, month in enumerate(months, start=2):
            ws.cell(row=row, column=i, value=_tsek(sign * values[month]))
        ws.cell(row=row, column=source_col, value=source)
        return row

    def total(label: str, start: int, end: int) -> int:
        row = ws.max_row + 1
        ws.cell(row=row, column=1, value=label)
        for i in range(2, source_col):
            ws.cell(row=row, column=i, value=f"=SUM({_col(i)}{start}:{_col(i)}{end})")
        ws.cell(row=row, column=source_col, value=f"Formel: summa rad {start} till {end}")
        _bold(ws, row)
        return row

    asset_rows = [
        write(sl.line, sl.values, 1, sl.source_ref) for sl in lines if sl.line in asset_lines
    ]
    assets = total("Summa tillgångar", asset_rows[0], asset_rows[-1])
    claim_rows = [
        write(sl.line, sl.values, -1, sl.source_ref) for sl in lines if sl.line not in asset_lines
    ]
    claims = total("Summa eget kapital och skulder", claim_rows[0], claim_rows[-1])
    row = ws.max_row + 1
    ws.cell(row=row, column=1, value="Differens")
    for i in range(2, source_col):
        ws.cell(row=row, column=i, value=f"={_col(i)}{assets}-{_col(i)}{claims}")
    ws.cell(
        row=row,
        column=source_col,
        value=f"Formel: tillgångar minus eget kapital och skulder, ska vara 0. "
        f"{COMPUTED_RESULT} gör att balansräkningen går ihop före bokslut.",
    )
    _bold(ws, row)
    _format(ws, range(2, source_col), range(2, ws.max_row + 1), TSEK)
    ws.freeze_panes = "B2"
    _widths(ws, {1: 44})


def _refs(refs: list[str]) -> str:
    return ", ".join(refs)


def _bridge(
    ws: Worksheet, result: AnalysisResult, rr_rows: dict[str, int], period_col: str
) -> tuple[int, int]:
    """Write EBITDA brygga; return the rows of Justerad EBITDA and Justerad nettoomsättning."""
    bridge = result.bridge
    _header(
        ws,
        [
            "Post",
            "Kategori",
            "Konton",
            "Verifikat",
            "Belopp TSEK",
            "Inkludera",
            "Varav nettoomsättning TSEK",
            "Not",
            "Källa",
        ],
    )
    basis = "LTM" if bridge.basis == "LTM" else "hittills i år"
    ws.append(
        [
            f"Rapporterad EBITDA ({basis}, {bridge.period[0]} till {bridge.period[1]})",
            None,
            None,
            None,
            f"='RR månad'!{period_col}{rr_rows['EBITDA']}",
            None,
            None,
            None,
            "RR månad, raden EBITDA, periodkolumnen",
        ]
    )
    items: list[BridgeItem] = [*bridge.items, *bridge.candidates]
    for item in items:
        ws.append(
            [
                item.label,
                CATEGORY_LABELS[item.category],
                ", ".join(item.accounts),
                _refs(item.voucher_refs),
                _tsek(item.amount),
                1 if item.included else 0,
                _tsek(item.revenue_amount),
                item.note or "",
                item.source_ref,
            ]
        )
    adjusted = ws.max_row + 1
    formula = f"=E2+SUMPRODUCT(E3:E{adjusted - 1},F3:F{adjusted - 1})" if items else "=E2"
    ws.append(
        [
            "Justerad EBITDA",
            None,
            None,
            None,
            formula,
            None,
            None,
            "Sätt Inkludera till 1 eller 0 för att ta med eller utesluta en post.",
            "Formel: rapporterad EBITDA plus belopp gånger inkludera",
        ]
    )
    _bold(ws, adjusted)
    period = f"{bridge.period[0]} till {bridge.period[1]}"
    ws.append(
        [
            f"Rapporterad nettoomsättning ({basis}, {period})",
            None,
            None,
            None,
            f"='RR månad'!{period_col}{rr_rows['Nettoomsättning']}",
            None,
            None,
            None,
            "RR månad, raden Nettoomsättning, periodkolumnen",
        ]
    )
    revenue = ws.max_row
    ws.append(
        [
            "Justerad nettoomsättning",
            None,
            None,
            None,
            f"=E{revenue}+SUMPRODUCT(G3:G{adjusted - 1},F3:F{adjusted - 1})"
            if items
            else f"=E{revenue}",
            None,
            None,
            "Bas för justerad EBITDA-marginal; poster på konto 3000 till 3699 följer Inkludera.",
            "Formel: rapporterad nettoomsättning plus varav nettoomsättning gånger inkludera",
        ]
    )
    adjusted_revenue = ws.max_row
    _bold(ws, adjusted_revenue)
    _format(ws, [5, 7], range(2, adjusted_revenue + 1), TSEK)
    ws.freeze_panes = "A2"
    _widths(ws, {1: 50, 4: 40, 7: 16, 8: 50, 9: 60})
    return adjusted, adjusted_revenue


def _working_capital(ws: Worksheet, result: AnalysisResult) -> tuple[int, int | None]:
    months = result.statements.months
    _header(ws, ["Rad", *months, "Källa"])
    source_col = len(months) + 2

    rows = []
    for line, members in nwc_lines(result.statements, result.mapping):
        row = ws.max_row + 1
        rows.append(row)
        ws.cell(row=row, column=1, value=NWC_LABELS.get(line, line))
        for i, month in enumerate(months, start=2):
            value = sum((result.statements.account_balances[a][month] for a in members), Decimal(0))
            ws.cell(row=row, column=i, value=_tsek(value))
        ws.cell(
            row=row,
            column=source_col,
            value=f"SIE #IB plus #TRANS vid månadens slut, konton {', '.join(members)}",
        )
    nwc = ws.max_row + 1
    ws.cell(row=nwc, column=1, value="Rörelsekapital")
    for i in range(2, source_col):
        total = f"=SUM({_col(i)}{rows[0]}:{_col(i)}{rows[-1]})" if rows else 0
        ws.cell(row=nwc, column=i, value=total)
    ws.cell(row=nwc, column=source_col, value=result.kpis.nwc.source_ref)
    _bold(ws, nwc)

    normal: int | None = None
    if len(months) >= 12:
        normal = nwc + 1
        last = len(months) + 1
        ws.cell(row=normal, column=1, value="Normalnivå, snitt 12 månader")
        ws.cell(
            row=normal,
            column=last,
            value=f"=AVERAGE({_col(last - 11)}{nwc}:{_col(last)}{nwc})",
        )
        ws.cell(row=normal, column=source_col, value=result.kpis.nwc_normal.source_ref)
        _bold(ws, normal)
    _format(ws, range(2, source_col), range(2, ws.max_row + 1), TSEK)
    ws.freeze_panes = "B2"
    _widths(ws, {1: 44})
    return nwc, normal


def _net_debt(ws: Worksheet, result: AnalysisResult) -> int:
    kpis = result.kpis
    _header(ws, ["Post", "Konton", "Belopp TSEK", "Källa"])
    for item in kpis.net_debt_items:
        ws.append([item.label, ", ".join(item.accounts), _tsek(item.amount), item.source_ref])
    total = ws.max_row + 1
    ws.append(
        [
            "Nettoskuld",
            None,
            f"=SUM(C2:C{total - 1})",
            f"Formel: summa av posterna per {kpis.net_debt_month}. Positivt belopp ökar "
            "nettoskulden, negativ nettoskuld är nettokassa.",
        ]
    )
    _bold(ws, total)
    if kpis.debt_like_items:
        ws.append([])
        ws.append(
            [
                "Att bedöma, ingår inte i nettoskulden",
                None,
                None,
                "Positivt belopp är en skuld, negativt en fordran",
            ]
        )
        _bold(ws, ws.max_row)
        for item in kpis.debt_like_items:
            ws.append([item.label, ", ".join(item.accounts), _tsek(item.amount), item.source_ref])
    _format(ws, [3], range(2, ws.max_row + 1), TSEK)
    ws.freeze_panes = "A2"
    _widths(ws, {1: 50, 4: 70})
    return total


def flag_statuses(result: AnalysisResult) -> dict[str, str]:
    """Voucher ref to its status in the bridge; anything not listed is Normal."""
    bridge = result.bridge
    status: dict[str, str] = {}
    for flags, label in (
        (bridge.outside_period, "Utanför perioden"),
        (bridge.not_adjusted, "Ej justerad, fråga"),
        (bridge.review, "Kräver granskning"),
    ):
        status.update({f.voucher_ref: label for f in flags})
    status.update({r: "Justerad" for item in bridge.items for r in item.voucher_refs})
    return status


def _flagged(ws: Worksheet, result: AnalysisResult) -> None:
    header = [
        "Verifikat",
        "Datum",
        "Konto",
        "Kontonamn",
        "Text",
        "Belopp i resultatet TSEK",
        "EBITDA-justering TSEK",
        "Kategori",
        "Konfidens",
        "Status",
        "Motivering",
        "Förfilter",
        "Källa",
    ]
    _header(ws, header)
    reasons = {h.voucher_ref: "; ".join(h.reasons) for h in result.prefilter_hits}
    statuses = flag_statuses(result)
    rows = sorted(
        ((flag, statuses.get(flag.voucher_ref, "Normal")) for flag in result.flags),
        key=lambda fs: (STATUS_ORDER.index(fs[1]), -abs(fs[0].amount)),
    )
    for flag, status in rows:
        ws.append(
            [
                flag.voucher_ref,
                flag.voucher_date,
                flag.account,
                flag.account_name,
                flag.voucher_text,
                _tsek(-flag.amount),
                _tsek(flag.ebitda_adjustment),
                CATEGORY_LABELS[flag.category],
                flag.confidence,
                status,
                flag.rationale,
                reasons.get(flag.voucher_ref, ""),
                f"SIE verifikat {flag.voucher_ref}, konto {flag.account}",
            ]
        )
    end = ws.max_row + 1
    _format(ws, [6, 7], range(2, end), TSEK)
    _format(ws, [9], range(2, end), CONFIDENCE)
    _format(ws, [2], range(2, end), "yyyy-mm-dd")
    ws.freeze_panes = "A2"
    _widths(ws, {5: 45, 11: 60, 12: 60, 13: 30})


def _questions(ws: Worksheet, result: AnalysisResult) -> None:
    _header(ws, ["Nr", "Fråga", "Konto", "Verifikat", "Belopp i resultatet TSEK", "Källa"])
    for i, q in enumerate(result.questions, start=1):
        ws.append(
            [
                i,
                q.text,
                ", ".join(q.accounts),
                _refs(q.voucher_refs),
                _tsek(-q.amount) if q.amount is not None else None,
                "Flaggade poster, verifikatens EBITDA-konton (intäkt positiv, kostnad negativ)",
            ]
        )
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    _format(ws, [5], range(2, ws.max_row + 1), TSEK)
    ws.freeze_panes = "A2"
    _widths(ws, {1: 6, 2: 90, 4: 40, 6: 45})


def _reconciliation(ws: Worksheet, result: AnalysisResult) -> int:
    """Write Avstämning; return the row of the status cell."""
    rec = result.reconciliation
    years = {fy.index: f"{fy.start} till {fy.end}" for fy in result.statements.fiscal_years}
    _header(
        ws,
        [
            "Räkenskapsår",
            "Mot",
            "Konto",
            "Beräknat SEK",
            "Enligt fil SEK",
            "Differens SEK",
            "Absolut differens SEK",
            "Källa",
        ],
    )
    for r in rec.rows:
        row = ws.max_row + 1
        ws.append(
            [
                years.get(r.year_index, str(r.year_index)),
                r.against,
                r.account,
                r.computed,
                r.stated,
                f"=D{row}-E{row}",
                f"=ABS(F{row})",
                f"Månadsrapporter mot SIE #{r.against} {r.year_index}, konto {r.account}",
            ]
        )
    end = ws.max_row
    # Fixed positions below the table, one blank row in between. (Appending an empty row
    # does not advance openpyxl's max_row, so rows are addressed explicitly.)
    total, status, unmapped, outside = end + 2, end + 3, end + 4, end + 5

    def put(row: int, cells: dict[int, Any]) -> None:
        for column, value in cells.items():
            ws.cell(row=row, column=column, value=value)

    put(
        total,
        {
            1: "Summa absolut differens",
            7: f"=SUM(G2:G{end})" if end >= 2 else 0,
            8: "Formel: summa av absoluta differenser",
        },
    )
    put(
        status,
        {
            1: "Status",
            2: f'=IF(G{total}+C{unmapped}+C{outside}=0,"OK","AVVIKER")',
            8: "Formel: OK när differenserna är 0 och inget konto eller verifikat faller utanför",
        },
    )
    put(
        unmapped,
        {
            1: "Omappade konton",
            2: ", ".join(rec.unmapped_accounts) or "inga",
            3: len(rec.unmapped_accounts),
            8: "Konton utanför BAS-intervallen, antal i kolumn C",
        },
    )
    put(
        outside,
        {
            1: "Verifikat utanför räkenskapsåren",
            2: ", ".join(rec.vouchers_outside_years) or "inga",
            3: len(rec.vouchers_outside_years),
            8: "SIE #VER jämfört med #RAR, antal i kolumn C",
        },
    )
    _bold(ws, status)
    _format(ws, [4, 5, 6, 7], range(2, total + 1), SEK)
    ws.freeze_panes = "A2"
    _widths(ws, {8: 60})
    return status


def _summary(
    ws: Worksheet,
    result: AnalysisResult,
    rr_rows: dict[str, int],
    period_col: str,
    adjusted_rows: tuple[int, int],
    nwc: tuple[int, int | None],
    net_debt_row: int,
) -> None:
    kpis, bridge = result.kpis, result.bridge
    basis = "LTM" if bridge.basis == "LTM" else "hittills i år"
    adjusted_row, adjusted_revenue_row = adjusted_rows
    months = result.statements.months
    last_col = _col(len(months) + 1)
    _header(ws, ["Uppgift", "Värde", "Källa"])
    # Rows are appended in order; formulas refer to earlier rows through `at`.
    at: dict[str, int] = {}

    def add(key: str, label: str, value: Any, source: str, number_format: str | None) -> None:
        ws.append([label, value, source])
        at[key] = ws.max_row
        if number_format:
            ws.cell(row=ws.max_row, column=2).number_format = number_format

    add("bolag", "Bolag", result.company_name, "SIE #FNAMN", None)
    add("orgnr", "Organisationsnummer", result.org_nr, "SIE #ORGNR", None)
    add("fil", "Källfil", result.source, "SIE4-fil", None)
    add("period", "Analyserad period", f"{months[0]} till {months[-1]}", "Verifikationsdatum", None)
    add(
        "ai",
        "AI-stöd",
        result.model or "Av (--no-llm), flaggor kräver granskning",
        "DILIGENS_MODEL",
        None,
    )
    rr = "'RR månad'!"
    add(
        "oms",
        f"Nettoomsättning {basis}",
        f"={rr}{period_col}{rr_rows['Nettoomsättning']}",
        "RR månad",
        TSEK,
    )
    add("ebitda", f"EBITDA {basis}", f"={rr}{period_col}{rr_rows['EBITDA']}", "RR månad", TSEK)
    add(
        "marginal",
        f"EBITDA-marginal {basis}",
        f'=IFERROR(B{at["ebitda"]}/B{at["oms"]},"")',
        "Formel: EBITDA delat med nettoomsättning",
        PERCENT,
    )
    add(
        "justerad",
        f"Justerad EBITDA {basis}",
        f"='EBITDA brygga'!E{adjusted_row}",
        "EBITDA brygga, med valda Inkludera",
        TSEK,
    )
    add(
        "justerad_oms",
        f"Justerad nettoomsättning {basis}",
        f"='EBITDA brygga'!E{adjusted_revenue_row}",
        "EBITDA brygga, intäktsposter med valda Inkludera",
        TSEK,
    )
    add(
        "justerad_marginal",
        f"Justerad EBITDA-marginal {basis}",
        f'=IFERROR(B{at["justerad"]}/B{at["justerad_oms"]},"")',
        "Formel: justerad EBITDA delat med justerad nettoomsättning",
        PERCENT,
    )
    if len(months) >= 24:
        n = len(months) + 1
        revenue = rr_rows["Nettoomsättning"]
        add(
            "tillvaxt",
            "Tillväxt LTM",
            f"=IFERROR(SUM({rr}{_col(n - 11)}{revenue}:{_col(n)}{revenue})"
            f'/SUM({rr}{_col(n - 23)}{revenue}:{_col(n - 12)}{revenue})-1,"")',
            "Formel: LTM mot LTM tolv månader tidigare",
            PERCENT,
        )
    add(
        "nwc",
        f"Rörelsekapital {months[-1]}",
        f"='Rörelsekapital'!{last_col}{nwc[0]}",
        "Rörelsekapital",
        TSEK,
    )
    if nwc[1] is not None:
        add(
            "nwc_normal",
            "Rörelsekapital normalnivå",
            f"='Rörelsekapital'!{last_col}{nwc[1]}",
            "Rörelsekapital, snitt 12 månader",
            TSEK,
        )
    add(
        "nettoskuld",
        f"Nettoskuld {kpis.net_debt_month}",
        f"='Nettoskuld'!C{net_debt_row}",
        "Nettoskuld, negativt värde är nettokassa",
        TSEK,
    )
    add(
        "avstamning",
        "Avstämning",
        "OK" if result.reconciliation.ok else "AVVIKER",
        "Avstämning mot SIE #RES och #UB, omappade konton och verifikat utanför åren",
        None,
    )
    add("flaggor", "Flaggade poster", len(result.flags), "Flaggade poster", None)
    add("granskning", "Kräver granskning", len(bridge.review), "Flaggade poster, status", None)
    add("fragor", "DD-frågor", len(result.questions), "DD frågor", None)
    ws.freeze_panes = "A2"
    _widths(ws, {1: 34, 2: 40, 3: 60})


def _sources(ws: Worksheet, result: AnalysisResult) -> None:
    threshold = f"{REVIEW_CONFIDENCE:.1f}".replace(".", ",")
    _header(ws, ["Uppgift", "Värde", "Källa"])
    rows = [
        ("Källfil", result.source, "SIE4-fil från bolagets bokföring"),
        ("Bolag", f"{result.company_name} ({result.org_nr})", "SIE #FNAMN och #ORGNR"),
        ("Genererad", dt.date.today().isoformat(), "Diligens"),
        ("AI-modell", result.model or "ingen (--no-llm)", "Inställning DILIGENS_MODEL"),
        ("Belopp", "TSEK, beräknade i kronor och delade med 1 000", "Beräknat i kod"),
        (
            "Tecken",
            "RR: intäkter positiva, kostnader negativa. BR: tillgångar, eget kapital och "
            "skulder positiva. Rörelsekapital: tillgångar positiva, skulder negativa.",
            "Presentationsregel",
        ),
        (
            "Konfidensgräns",
            f"Flaggor under {threshold} justerar inte EBITDA och listas som Kräver granskning",
            "Bedömningsregel",
        ),
        (
            "Siffror",
            "Alla belopp och nyckeltal är beräknade i kod. AI väljer bara kategori, motivering "
            "och frågetext.",
            "Diligens",
        ),
        (
            "Spårbarhet",
            "Varje flaggad post har verifikat och konto, varje nyckeltal sin källa",
            "Kolumnen Källa",
        ),
    ]
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "A2"
    _widths(ws, {1: 20, 2: 90, 3: 40})


def export_dd_pack(result: AnalysisResult, out_dir: Path) -> Path:
    """Write `dd_pack.xlsx` to `out_dir` and return its path."""
    wb = Workbook()
    sheets = {name: wb.create_sheet(name) for name in SHEETS}
    wb.remove(wb.worksheets[0])  # the default empty sheet

    rr_rows, period_col = _income_statement(sheets["RR månad"], result)
    _balance_sheet(sheets["BR månad"], result)
    adjusted_rows = _bridge(sheets["EBITDA brygga"], result, rr_rows, period_col)
    nwc = _working_capital(sheets["Rörelsekapital"], result)
    net_debt_row = _net_debt(sheets["Nettoskuld"], result)
    _flagged(sheets["Flaggade poster"], result)
    _questions(sheets["DD frågor"], result)
    _reconciliation(sheets["Avstämning"], result)
    _summary(
        sheets["Sammanfattning"], result, rr_rows, period_col, adjusted_rows, nwc, net_debt_row
    )
    _sources(sheets["Källor"], result)
    # openpyxl stores no computed values; make Excel calculate on open (also in Protected View).
    wb.calculation.fullCalcOnLoad = True

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "dd_pack.xlsx"
    wb.save(path)
    return path
