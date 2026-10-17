"""Buy & Build long list as an Excel workbook: one row per candidate with fit score, reasons
and the source of every figure (SCB field or iXBRL element and document)."""

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from diligens.models import Candidate, Figure, ScreeningResult
from diligens.screening import with_ranks

TSEK = "#,##0.0"
PERCENT = "0.0%"
THOUSAND = Decimal(1000)

COLUMNS = [
    "Rang",
    "Poäng",
    "Profil",
    "Nyckeltal",
    "Bolag",
    "Org.nr",
    "Ort",
    "Län",
    "SNI",
    "Omsättningsklass",
    "Anställdaklass",
    "Bokslut",
    "Nettoomsättning TSEK",
    "Föregående år TSEK",
    "Tillväxt",
    "Rörelseresultat TSEK",
    "Rörelsemarginal",
    "Skäl",
    "Källa",
]


def _value(figure: Figure | None, scale: Decimal = Decimal(1)) -> Decimal | None:
    return None if figure is None or figure.value is None else figure.value / scale


def _org(org_nr: str) -> str:
    return f"{org_nr[:6]}-{org_nr[6:]}" if len(org_nr) == 10 else org_nr


def _header(ws: Worksheet, labels: list[str]) -> None:
    ws.append(labels)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE3EA")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def _sources(candidate: Candidate, result: ScreeningResult) -> str:
    sni = candidate.unit.primary_sni
    parts = [
        f"SCB AFR juridiskaenheter/naringsgren/{sni}: namn, ort, län, SNI, anställdaklass",
    ]
    if candidate.turnover_class:
        parts.append(
            f"SCB AFR juridiskaenheter/omsattningsklass/{candidate.turnover_class}: "
            "omsättningsklass"
        )
    financial = candidate.financial
    if financial is not None:
        for label, figure in (
            ("Nettoomsättning", financial.revenue),
            ("Föregående år", financial.revenue_prior),
            ("Rörelseresultat", financial.ebit),
        ):
            if figure is not None:
                parts.append(f"{label}: {figure.source_ref}")
    return "; ".join(parts)


def export_long_list(result: ScreeningResult, out_dir: Path) -> Path:
    """Write `long_list.xlsx` to `out_dir` and return its path."""
    labels = result.labels
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Long list"
    _header(ws, COLUMNS)
    for rank, c in with_ranks(result.candidates):
        row = ws.max_row + 1
        f = c.financial
        reasons = c.profile_reasons + (f.reasons if f else [])
        ws.append(
            [
                rank,
                f"=C{row}+D{row}",
                c.profile_points,
                f.points if f else None,
                c.unit.name,
                _org(c.org_nr),
                c.unit.post_town.title(),
                labels.get("lankoder", {}).get(c.unit.county, c.unit.county),
                c.unit.primary_sni,
                labels.get("omsklkoder", {}).get(c.turnover_class, c.turnover_class or "okänd"),
                labels.get("anstklkoder", {}).get(c.unit.employee_class, c.unit.employee_class),
                f.status if f else "Ej hämtad",
                _value(f.revenue, THOUSAND) if f else None,
                _value(f.revenue_prior, THOUSAND) if f else None,
                _value(f.growth) if f else None,
                _value(f.ebit, THOUSAND) if f else None,
                _value(f.ebit_margin) if f else None,
                "; ".join(reasons),
                _sources(c, result),
            ]
        )
        for column in (13, 14, 16):
            ws.cell(row=row, column=column).number_format = TSEK
        for column in (15, 17):
            ws.cell(row=row, column=column).number_format = PERCENT
    widths = {
        1: 6,
        2: 8,
        3: 8,
        4: 10,
        5: 34,
        6: 13,
        7: 16,
        8: 16,
        9: 8,
        10: 20,
        11: 16,
        12: 18,
        13: 14,
        14: 14,
        15: 10,
        16: 14,
        17: 12,
        18: 70,
        19: 90,
    }
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width

    criteria = result.criteria
    sheet = wb.create_sheet("Kriterier")
    _header(sheet, ["Kriterium", "Värde", "Källa"])

    def describe(codes: list[str], table: str) -> str:
        names = labels.get(table, {})
        return ", ".join(f"{code} {names.get(code, '')}".strip() for code in codes) or "alla"

    for name, value, source in (
        ("SNI (primär näringsgren)", ", ".join(criteria.sni), "SCB naringsgrenkoder, SNI 2025"),
        (
            "Omsättningsklass",
            describe(criteria.turnover_classes, "omsklkoder"),
            "SCB omsklkoder; en klass ifrån ger halva poängen",
        ),
        (
            "Anställdaklass",
            describe(criteria.employee_classes, "anstklkoder"),
            "SCB anstklkoder; en klass ifrån ger halva poängen",
        ),
        ("Län", describe(criteria.counties, "lankoder"), "SCB lankoder, säte"),
        ("Bolag med SNI hos SCB", str(result.units_scanned), "SCB AFR"),
        ("Kandidater efter filter", str(len(result.candidates)), "Verksamma aktiebolag"),
        ("Årsredovisningar sökta", str(result.reports_fetched), "Bolagsverket, bästa profilerna"),
        ("Genererad", result.generated.isoformat(), "Diligens"),
        (
            "Poäng",
            "Profil 0 till 60 (SNI 15, omsättning 15, anställda 10, län 10, verksamt "
            "aktiebolag 10) och nyckeltal 0 till 40 (tillväxt 20, rörelsemarginal 20). SNI och "
            "verksamt aktiebolag är krav, så alla kandidater har de 25 poängen; det är "
            "omsättning, anställda, län och nyckeltal som skiljer dem åt.",
            "Poängregel",
        ),
        (
            "Rang",
            "Bara bolag med nyckeltalspoäng rangordnas, så att varje rang jämför samma slags "
            "poäng. Sökta bolag utan nyckeltal (saknas digitalt, kunde inte läsas eller "
            "hämtas) följer utan rang i profilordning, sist de som inte söktes. Tomt i "
            "Nyckeltal betyder att poäng saknas, inte noll.",
            "Poängregel",
        ),
        *(("Varning", note, "Körningen") for note in result.notes),
    ):
        sheet.append([name, value, source])
    for column, width in {1: 28, 2: 70, 3: 50}.items():
        sheet.column_dimensions[get_column_letter(column)].width = width

    # openpyxl stores no computed values; make Excel calculate on open (also Protected View).
    wb.calculation.fullCalcOnLoad = True
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "long_list.xlsx"
    wb.save(path)
    return path
