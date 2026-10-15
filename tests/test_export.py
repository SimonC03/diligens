"""M4 export tests. One test per acceptance criterion in docs/PLAN.md. No API calls."""

import json
import re
from pathlib import Path
from types import SimpleNamespace
from typing import Any, cast

import pytest
from anthropic import Anthropic
from formula_eval import FormulaEvaluator
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from typer.testing import CliRunner

from diligens.agent.memo import memo_facts, write_memo
from diligens.cli import app
from diligens.config import Settings
from diligens.export.excel import SHEETS, export_dd_pack
from diligens.models import AnalysisResult, Facit
from diligens.pipeline import run_analysis
from diligens.sie.generator import generate_demo_company

# Category and confidence per planted anomaly, as a model might answer. The owner fee comes
# back as related party (a candidate, not adjusted) and the car below the threshold.
ANSWERS = {
    1: ("narstaende", 0.9),
    2: ("engangspost", 0.95),
    3: ("agarrelaterad", 0.6),
    4: ("periodisering", 0.85),
    5: ("engangspost", 0.9),
    6: ("engangspost", 0.9),
}


class ScriptedMessages:
    def __init__(self, facit: Facit) -> None:
        self.answer = {ref: ANSWERS[a.id] for a in facit.anomalies for ref in a.voucher_refs}
        self.calls = 0

    def create(self, **kwargs: Any) -> SimpleNamespace:
        self.calls += 1
        content = kwargs["messages"][0]["content"]
        refs = re.findall(r'"ref": "(\w+)"', content)
        ids = re.findall(r'"id": "(\w+)"', content)
        if refs:
            output: dict[str, Any] = {
                "flags": [
                    {
                        "voucher_ref": r,
                        "category": self.answer.get(r, ("normal", 0.9))[0],
                        "rationale": "Bedömning enligt underlaget.",
                        "confidence": self.answer.get(r, ("normal", 0.9))[1],
                    }
                    for r in refs
                ]
            }
        elif ids:
            output = {"questions": [{"id": i, "text": "Vad avser posterna?"} for i in ids]}
        else:
            output = {"markdown": "## Sammanfattning\n\nEBITDA uppgår till {ebitda}."}
        return SimpleNamespace(
            content=[SimpleNamespace(type="text", text=json.dumps(output, ensure_ascii=False))],
            stop_reason="end_turn",
            usage=SimpleNamespace(input_tokens=10, output_tokens=10, cache_read_input_tokens=0),
        )


@pytest.fixture(scope="module")
def analysed(tmp_path_factory: pytest.TempPathFactory) -> tuple[AnalysisResult, Path]:
    folder = tmp_path_factory.mktemp("export")
    sie = folder / "demo_bolag.se"
    facit = generate_demo_company(42, sie)
    client = SimpleNamespace(messages=ScriptedMessages(facit))
    settings = Settings(_env_file=None, anthropic_api_key=None, diligens_model="claude-test")  # type: ignore[call-arg]
    result = run_analysis(sie, use_llm=True, client=cast(Anthropic, client), settings=settings)
    return result, export_dd_pack(result, folder / "out")


def _wb(path: Path) -> Workbook:
    return load_workbook(path)


def _row_of(ws: Any, label: str, column: int = 1) -> int:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=column).value == label:
            return row
    raise AssertionError(f"{label!r} not found in {ws.title}")


def _header(ws: Any) -> list[Any]:
    return [c.value for c in ws[1]]


def test_workbook_has_all_sheets_with_sources(analysed: tuple[AnalysisResult, Path]) -> None:
    _, path = analysed
    wb = _wb(path)

    assert wb.sheetnames == [
        "Sammanfattning",
        "RR månad",
        "BR månad",
        "EBITDA brygga",
        "Rörelsekapital",
        "Nettoskuld",
        "Flaggade poster",
        "DD frågor",
        "Avstämning",
        "Källor",
    ]
    assert wb.sheetnames == SHEETS
    for ws in wb.worksheets:
        assert "Källa" in _header(ws), ws.title
        assert ws.max_row > 2, ws.title


def test_totals_and_bridge_are_live_formulas(analysed: tuple[AnalysisResult, Path]) -> None:
    result, path = analysed
    wb = _wb(path)
    ev = FormulaEvaluator(wb)
    kpis, bridge = result.kpis, result.bridge
    thousand = 1000

    rr = wb["RR månad"]
    ebitda_row = _row_of(rr, "EBITDA")
    months = _header(rr)
    for month in ("2024-07", "2025-12"):
        cell = rr.cell(row=ebitda_row, column=months.index(month) + 1)
        assert str(cell.value).startswith("=")
        expected = kpis.monthly["ebitda"].values[month]
        assert ev.value("RR månad", cell.coordinate) == pytest.approx(float(expected) / thousand)

    br = wb["BR månad"]
    diff_row = _row_of(br, "Differens")
    for column in range(2, len(_header(br))):
        cell = br.cell(row=diff_row, column=column)
        assert str(cell.value).startswith("=")
        assert ev.value("BR månad", cell.coordinate) == pytest.approx(0, abs=1e-6)

    ws = wb["EBITDA brygga"]
    reported = ws.cell(row=_row_of(ws, "Justerad EBITDA"), column=5)
    assert str(reported.value).startswith("=")
    assert "SUMPRODUCT" in reported.value
    assert ev.value(ws.title, reported.coordinate) == pytest.approx(
        float(bridge.adjusted) / thousand
    )
    start = ws.cell(row=2, column=5)
    assert ev.value(ws.title, start.coordinate) == pytest.approx(float(bridge.reported) / thousand)

    # The related party fee is listed with Inkludera 0; switching it on moves the total.
    assert bridge.candidates and not any(c.included for c in bridge.candidates)
    related = next(c for c in bridge.candidates if c.category == "narstaende")
    row = _row_of(ws, related.label)
    assert ws.cell(row=row, column=6).value == 0
    ws.cell(row=row, column=6).value = 1
    expected = (bridge.adjusted + related.amount) / thousand
    assert ev.value(ws.title, reported.coordinate) == pytest.approx(float(expected))

    nwc = wb["Rörelsekapital"]
    header = _header(nwc)
    last = nwc.cell(row=_row_of(nwc, "Rörelsekapital"), column=header.index("2025-12") + 1)
    assert ev.value(nwc.title, last.coordinate) == pytest.approx(
        float(kpis.nwc.values["2025-12"] or 0) / thousand
    )
    normal = nwc.cell(row=_row_of(nwc, "Normalnivå, snitt 12 månader"), column=last.column)
    assert "AVERAGE" in normal.value
    assert ev.value(nwc.title, normal.coordinate) == pytest.approx(
        float(kpis.nwc_normal.value or 0) / thousand
    )

    debt = wb["Nettoskuld"]
    total = debt.cell(row=_row_of(debt, "Nettoskuld"), column=3)
    assert str(total.value).startswith("=SUM")
    assert ev.value(debt.title, total.coordinate) == pytest.approx(float(kpis.net_debt) / thousand)

    summary = wb["Sammanfattning"]
    linked = summary.cell(row=_row_of(summary, "Justerad EBITDA"), column=2)
    assert "'EBITDA brygga'!" in linked.value


def test_tsek_format_frozen_headers_and_swedish_labels(
    analysed: tuple[AnalysisResult, Path],
) -> None:
    _, path = analysed
    wb = _wb(path)

    rr = wb["RR månad"]
    assert rr.freeze_panes == "B2"
    revenue = rr.cell(row=_row_of(rr, "Nettoomsättning"), column=2)
    assert revenue.number_format == "#,##0.0" and revenue.value > 0  # TSEK, revenue positive
    costs = rr.cell(row=_row_of(rr, "Personalkostnader"), column=2)
    assert costs.value < 0  # costs negative
    assert _header(rr)[0] == "Rad" and _header(rr)[-1] == "Källa"
    assert rr.column_dimensions["A"].width >= 30

    for ws in wb.worksheets:
        assert ws.freeze_panes in ("A2", "B2"), ws.title
        assert all(isinstance(v, str) for v in _header(ws) if v is not None), ws.title
    flagged = wb["Flaggade poster"]
    assert _header(flagged)[:4] == ["Verifikat", "Datum", "Konto", "Kontonamn"]
    assert flagged.cell(row=2, column=_header(flagged).index("Konfidens") + 1).number_format == (
        "0.00"
    )
    summary = wb["Sammanfattning"]
    margin = summary.cell(row=_row_of(summary, "EBITDA-marginal LTM"), column=2)
    assert margin.number_format == "0.0%"


def test_every_flagged_row_shows_voucher_and_account(
    analysed: tuple[AnalysisResult, Path],
) -> None:
    result, path = analysed
    ws = _wb(path)["Flaggade poster"]
    header = _header(ws)
    col = {name: header.index(name) + 1 for name in header}

    rows = [
        {name: ws.cell(row=r, column=col[name]).value for name in header}
        for r in range(2, ws.max_row + 1)
    ]
    assert {r["Verifikat"] for r in rows} == {f.voucher_ref for f in result.flags}
    assert all(r["Verifikat"] and r["Konto"] and r["Källa"] for r in rows)
    assert all(r["Verifikat"] in r["Källa"] and str(r["Konto"]) in r["Källa"] for r in rows)
    statuses = {r["Status"] for r in rows}
    assert {"Justerad", "Kräver granskning", "Normal"} <= statuses
    late = next(r for r in rows if r["Kategori"] == "Periodisering")
    assert late["EBITDA-justering TSEK"] == pytest.approx(-600.0)


def test_cli_writes_workbook_that_opens_with_formulas(tmp_path: Path) -> None:
    sie = tmp_path / "demo_bolag.se"
    generate_demo_company(42, sie)
    out = tmp_path / "output"

    run = CliRunner().invoke(app, ["analyze", str(sie), "--out", str(out), "--no-llm"])

    assert run.exit_code == 0, run.output
    assert (out / "dd_pack.xlsx").exists() and (out / "memo.md").exists()
    assert "dd_pack.xlsx" in run.output
    wb = load_workbook(out / "dd_pack.xlsx")
    assert wb.sheetnames == SHEETS
    formulas = sum(
        1
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    assert formulas > 100


_NUMBER = re.compile(r"\d[\d\s]*(?:,\d+)?")


def _numbers(text: str) -> set[str]:
    return {" ".join(n.split()) for n in _NUMBER.findall(text)}


class MemoMessages:
    def __init__(self, replies: list[str]) -> None:
        self.replies = replies
        self.calls: list[dict[str, Any]] = []

    def create(self, **kwargs: Any) -> SimpleNamespace:
        self.calls.append(kwargs)
        text = self.replies[min(len(self.calls), len(self.replies)) - 1]
        return SimpleNamespace(
            content=[SimpleNamespace(type="text", text=json.dumps({"markdown": text}))],
            stop_reason="end_turn",
            usage=SimpleNamespace(input_tokens=10, output_tokens=10, cache_read_input_tokens=0),
        )


def test_memo_uses_computed_numbers_only(analysed: tuple[AnalysisResult, Path]) -> None:
    result, _ = analysed
    facts = memo_facts(result)
    allowed = {n for value in facts.values() for n in _numbers(value)}
    assert facts["ebitda"].endswith("TSEK") and "," in facts["ebitda"]

    invented = "## Sammanfattning\n\nEBITDA uppgick till 5 300 TSEK, en stark nivå."
    good = (
        "## Sammanfattning\n\n{bolag} redovisar en nettoomsättning om {omsattning} och en "
        "EBITDA om {ebitda} ({bas}). Justerad EBITDA uppgår till {justerad_ebitda}.\n\n"
        "## Frågor till säljaren\n\n- {fraga_1}\n- {fraga_2}"
    )
    messages = MemoMessages([invented, good])
    memo = write_memo(result, cast(Anthropic, SimpleNamespace(messages=messages)), model="x")

    assert len(messages.calls) == 2
    assert "siffror" in messages.calls[1]["messages"][0]["content"]
    assert facts["ebitda"] in memo and facts["justerad_ebitda"] in memo
    assert "{" not in memo
    assert _numbers(memo) <= allowed, _numbers(memo) - allowed

    # Unknown placeholders never get through; after the retries the template is used.
    unknown = MemoMessages(["## Sammanfattning\n\nMarginalen är {okand_nyckel}."])
    fallback = write_memo(result, cast(Anthropic, SimpleNamespace(messages=unknown)), model="x")
    template = write_memo(result, None)
    assert fallback == template and len(unknown.calls) == 3

    for text in (memo, template):
        assert len(text.split()) <= 450  # one page
        assert _numbers(text) <= allowed, _numbers(text) - allowed
    for heading in ("## Sammanfattning", "## Normaliseringar", "## Frågor till säljaren"):
        assert heading in template
    assert result.memo  # the pipeline fills it in


@pytest.mark.live
def test_live_memo(analysed: tuple[AnalysisResult, Path]) -> None:
    """Real API call that writes the memo. Run with: uv run pytest -m live"""
    settings = Settings()
    if settings.anthropic_api_key is None:
        pytest.skip("ANTHROPIC_API_KEY saknas")
    result, _ = analysed
    client = Anthropic(api_key=settings.anthropic_api_key.get_secret_value())

    memo = write_memo(result, client, model=settings.diligens_model)

    allowed = {n for value in memo_facts(result).values() for n in _numbers(value)}
    assert "## Sammanfattning" in memo and "{" not in memo
    assert _numbers(memo) <= allowed
    assert len(memo.split()) <= 450
