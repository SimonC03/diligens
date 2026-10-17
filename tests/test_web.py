"""M8 web app tests: pages, the analysis and screening flows, the advisor's choices and
local-only safety. Synthetic data, fake SCB and Bolagsverket, no network and no LLM."""

import datetime as dt
import io
import threading
import time
from collections.abc import Iterator
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from helpers import fake_registry, fake_scb
from helpers import unit as scb_unit
from openpyxl import load_workbook

from diligens.analysis.bridge import ebitda_bridge
from diligens.config import Settings
from diligens.models import Flag, FlagCategory, Ledger
from diligens.pipeline import run_analysis
from diligens.sie.generator import build_demo_company, generate_demo_company
from diligens.sources.bolagsverket import BolagsverketUnavailable
from diligens.web import views
from diligens.web.app import create_app
from diligens.web.jobs import Reporter
from diligens.web.workspace import AnalysisMeta, Workspace

SECRETS = ("scb-key-not-real", "bv-id-not-real", "bv-secret-not-real")
PAGES = (
    "/",
    "/analyser",
    "/analyser/ny",
    "/screening",
    "/screening/ny",
    "/metod",
    "/installningar",
)


def _settings() -> Settings:
    return Settings(  # type: ignore[call-arg]
        _env_file=None,
        anthropic_api_key=None,
        scb_api_key=SECRETS[0],
        bolagsverket_client_id=SECRETS[1],
        bolagsverket_client_secret=SECRETS[2],
    )


@pytest.fixture(scope="module")
def demo(tmp_path_factory: pytest.TempPathFactory) -> Path:
    path = tmp_path_factory.mktemp("demo") / "demo_bolag.se"
    generate_demo_company(42, path)
    return path


def _units() -> tuple[list[dict[str, object]], dict[str, list[str]]]:
    units = [scb_unit(1), scb_unit(2, lan="12"), scb_unit(3, jur="10"), scb_unit(9)]
    pe = {n: f"1655990{n:05d}" for n in (1, 2, 3, 9)}
    return units, {"6": [pe[1], pe[2], pe[3]], "7": [pe[9]]}


def _app(workspace: Path, demo: Path) -> TestClient:
    units, turnover = _units()
    app = create_app(
        settings_provider=_settings,
        workspace=workspace,
        scb_factory=lambda s: fake_scb(units, turnover),  # type: ignore[arg-type]
        bolagsverket_factory=lambda s: fake_registry(),
        inline_jobs=True,
        demo_file=demo,
        allowed_hosts=("testserver",),
    )
    return TestClient(app)


@pytest.fixture
def client(tmp_path: Path, demo: Path) -> Iterator[TestClient]:
    with _app(tmp_path / "arbetsyta", demo) as c:
        yield c


def test_every_page_renders_without_secrets(client: TestClient) -> None:
    for path in PAGES:
        page = client.get(path)
        assert page.status_code == 200, path
        assert "Diligens" in page.text and "Due diligence" in page.text
        assert not any(secret in page.text for secret in SECRETS), path
    assert "Koden räknar" in client.get("/metod").text
    missing = client.get("/analyser/20990101-000000-abcdef")
    assert missing.status_code == 404 and "finns inte" in missing.text


def test_analysis_flow_from_upload_to_excel(client: TestClient, tmp_path: Path, demo: Path) -> None:
    started = client.post(
        "/analyser",
        files={"fil": ("demo_bolag.se", demo.read_bytes(), "application/octet-stream")},
        data={"narstaende": "Ägarbolaget AB"},
    )
    # Upload, job and result: the inline job is done, so the redirects end on the summary.
    assert started.status_code == 200
    assert started.url.path.endswith("/sammanfattning")
    item_id = started.url.path.split("/")[2]
    assert "Demobolaget Konsult AB" in started.text and "Stämmer mot SIE-filen" in started.text

    result = run_analysis(demo, use_llm=False, related_parties=["Ägarbolaget AB"])
    assert views.tsek(result.bridge.adjusted) in started.text

    for tab in (
        "resultatrakning",
        "balansrakning",
        "rorelsekapital",
        "nettoskuld",
        "brygga",
        "flaggor",
        "fragor",
        "avstamning",
        "memo",
    ):
        page = client.get(f"/analyser/{item_id}/{tab}")
        assert page.status_code == 200, tab
    assert "Kräver granskning" in client.get(f"/analyser/{item_id}/flaggor").text
    assert client.get(f"/analyser/{item_id}/okand").status_code == 404

    # Every figure traces to a voucher: B473 as booked, as a page and as a drawer fragment.
    voucher = client.get(f"/analyser/{item_id}/verifikat/B473")
    assert voucher.status_code == 200 and "3041" in voucher.text and "1510" in voucher.text
    fragment = client.get(f"/analyser/{item_id}/verifikat/B473?del=1")
    assert "<html" not in fragment.text and "Verifikat B473" in fragment.text

    pack = client.get(f"/analyser/{item_id}/fil/dd_pack.xlsx")
    assert pack.status_code == 200
    assert load_workbook(io.BytesIO(pack.content)).sheetnames[0] == "Sammanfattning"
    assert client.get(f"/analyser/{item_id}/fil/memo.md").text.startswith("## Sammanfattning")

    # Saved in the workspace: a new app on the same folder lists it.
    with _app(tmp_path / "arbetsyta", demo) as again:
        assert "Demobolaget Konsult AB" in again.get("/analyser").text

    deleted = client.post(f"/analyser/{item_id}/ta-bort")
    assert deleted.status_code == 200 and "Inga analyser ännu" in deleted.text


def test_bad_uploads_give_readable_errors(client: TestClient) -> None:
    nothing = client.post("/analyser", data={})
    assert "Välj en SIE4-fil" in nothing.text

    wrong = client.post("/analyser", files={"fil": ("bild.png", b"\x89PNG not a SIE file")})
    assert wrong.status_code == 200 and "Körningen avbröts" in wrong.text
    assert "Inga analyser ännu" in client.get("/analyser").text  # nothing half saved


def test_statement_rows_agree_with_the_kpis(demo: Path) -> None:
    result = run_analysis(demo, use_llm=False)
    months = result.statements.months
    rows = {row.label: row for row in views.income_rows(result)}
    assert rows["EBITDA"].values == [result.kpis.monthly["ebitda"].values[m] for m in months]
    assert rows["EBITDA"].period == result.bridge.reported
    assert rows["Nettoomsättning"].period == result.bridge.revenue
    assert rows["EBIT"].values == [result.kpis.monthly["ebit"].values[m] for m in months]
    balance = {row.label: row for row in views.balance_rows(result)}
    assert all(v == 0 for v in balance["Differens"].values)
    nwc = {row.label: row for row in views.nwc_rows(result)}
    assert nwc["Rörelsekapital"].values == [result.kpis.nwc.values[m] for m in months]
    chart = views.monthly_chart(result)
    assert len(chart.bars) == len(months) and all(b.height >= 0 for b in chart.bars)
    assert views.memo_html("## Rubrik\n- <b>x</b> **fet**") == (
        "<h3>Rubrik</h3>\n<ul>\n<li>&lt;b&gt;x&lt;/b&gt; <strong>fet</strong></li>\n</ul>"
    )


def _flags(ledger: Ledger, refs: list[str], category: FlagCategory) -> list[Flag]:
    by_ref = {v.ref: v for v in ledger.vouchers}
    flags = []
    for ref in refs:
        main = max(
            (t for t in by_ref[ref].transactions if t.account[0] in "34567"),
            key=lambda t: abs(t.amount),
        )
        flags.append(
            Flag(
                voucher_ref=ref,
                account=main.account,
                amount=main.amount,
                category=category,
                rationale="Test.",
                confidence=0.9,
            )
        )
    return flags


def test_bridge_choice_recomputes_and_follows_into_excel(tmp_path: Path, demo: Path) -> None:
    # An analysis with confident flags on the planted anomalies, saved as the app would.
    ledger, facit = build_demo_company(42)
    a = {x.id: x for x in facit.anomalies}
    result = run_analysis(demo, use_llm=False)
    flags = _flags(ledger, a[2].voucher_refs, "engangspost") + _flags(
        ledger, a[4].voucher_refs, "periodisering"
    )
    result = result.model_copy(update={"bridge": ebitda_bridge(result.kpis, ledger, flags)})
    space = Workspace(tmp_path / "arbetsyta")
    item_id = space.new_id()
    space.upload_path(item_id).write_bytes(demo.read_bytes())
    space.save_analysis(
        AnalysisMeta(
            id=item_id,
            created=dt.datetime(2026, 10, 17, 9, 0),
            source_name="demo_bolag.se",
            company_name=result.company_name,
            org_nr=result.org_nr,
            period="2025-01 till 2025-12",
            basis="LTM",
            use_llm=True,
            model="test",
            reconciled=True,
        ),
        result,
    )
    items = result.bridge.items
    assert [i.category for i in items] == ["engangspost", "engangspost", "periodisering"]

    with _app(tmp_path / "arbetsyta", demo) as client:
        # Switch the revenue item off: the totals come back from with_selection.
        choice = [i.category != "periodisering" for i in items]
        answer = client.post(f"/api/analyser/{item_id}/urval", json={"inkludera": choice})
        assert answer.status_code == 200
        body = answer.json()
        expected = result.bridge.reported + sum(
            i.amount for i in items if i.category != "periodisering"
        )
        assert body["justerad_ebitda"] == views.tsek(expected)
        assert body["justerad_omsattning"] == views.tsek(result.bridge.revenue)
        assert body["andrat"] and not body["memo_aktuellt"]

        page = client.get(f"/analyser/{item_id}/brygga")
        assert views.tsek(expected) in page.text and "Urvalet har ändrats" in page.text

        pack = load_workbook(
            io.BytesIO(client.get(f"/analyser/{item_id}/fil/dd_pack.xlsx").content)
        )
        ws = pack["EBITDA brygga"]
        included = {row[1].value: row[5].value for row in ws.iter_rows(min_row=3) if row[1].value}
        assert included["Periodisering"] == 0 and included["Engångspost"] == 1
        # The flagged items follow the choice too: B473 is left out, with no adjustment.
        flagged = {row[0].value: row for row in pack["Flaggade poster"].iter_rows(min_row=2)}
        assert flagged["B473"][9].value == "Ej medräknad" and flagged["B473"][6].value == 0
        left_out = client.get(f"/analyser/{item_id}/flaggor", params={"status": "Ej medräknad"})
        assert "B473" in left_out.text and "D277" not in left_out.text

        # The memo is rewritten with the choice (template without AI) and is current again.
        client.post(f"/analyser/{item_id}/memo")
        assert "Urvalet i bryggan har ändrats" not in client.get(f"/analyser/{item_id}/memo").text
        memo = client.get(f"/analyser/{item_id}/fil/memo.md").text
        assert views.tsek(expected) in memo
        assert "EBITDA-effekt om den tas med -600,0 TSEK" in memo  # listed as left out
        assert "EBITDA-effekt -600,0" not in memo  # not among the included adjustments

        # Back to the proposal clears the choice; a wrong length is refused.
        reset = client.post(
            f"/api/analyser/{item_id}/urval", json={"inkludera": [True, True, True]}
        )
        assert not reset.json()["andrat"]
        assert not reset.json()["memo_aktuellt"]  # the memo was written for the choice
        again = client.post(f"/api/analyser/{item_id}/urval", json={"inkludera": choice})
        assert again.json()["memo_aktuellt"]
        bad = client.post(f"/api/analyser/{item_id}/urval", json={"inkludera": [True]})
        assert bad.status_code == 400 and "poster" in bad.json()["error"]


def test_screening_flow_with_form_tables_and_company_page(client: TestClient) -> None:
    form = client.get("/screening/ny")
    assert "62100 Dataprogrammering" in form.text and "Stockholm" in form.text

    done = client.post(
        "/screening",
        data={
            "sni": "62100",
            "oms_kl": ["6"],
            "anst_kl": ["5"],
            "lan": ["01"],
            "topp": "25",
            "med_bokslut": "1",
        },
    )
    assert done.status_code == 200 and "/screening/" in str(done.url)
    assert "Testbolag 1 AB" in done.text and "Saknas digitalt" in done.text
    assert "Testbolag 3 AB" not in done.text  # not an aktiebolag
    item_id = done.url.path.split("/")[2]

    company = client.get(f"/screening/{item_id}/bolag/5599000001")
    assert company.status_code == 200 and "Nettoomsättning" in company.text
    assert "iXBRL se-gen-base:Nettoomsattning" in company.text
    report = client.get(f"/screening/{item_id}/bolag/5599000001?bokslut=1")
    assert "Årsredovisning 2025-01-01 till 2025-12-31" in report.text
    assert "Soliditet" in report.text

    sheet = client.get(f"/screening/{item_id}/fil/long_list.xlsx")
    assert load_workbook(io.BytesIO(sheet.content)).sheetnames == ["Long list", "Kriterier"]
    assert "SNI 62100" in client.get("/").text  # listed under recent work

    bad = client.post("/screening", data={"sni": "6210"})
    assert "fem siffror" in bad.text


def test_local_only_and_same_site(client: TestClient) -> None:
    foreign_host = client.get("/", headers={"host": "evil.example"})
    assert foreign_host.status_code == 400
    for origin in (
        "https://evil.example",
        "http://localhost:3000",
        "http://testserver:8080",
        "null",
    ):
        refused = client.post("/analyser", data={"demo": "1"}, headers={"origin": origin})
        assert refused.status_code == 403, origin
    for site in ("cross-site", "same-site"):  # same-site: another port on this computer
        marked = client.post("/analyser", data={"demo": "1"}, headers={"sec-fetch-site": site})
        assert marked.status_code == 403, site
        assert marked.headers["content-type"].startswith("text/html; charset=utf-8")
        assert "Förfrågan kom från en annan webbplats" in marked.text

    # A real form post from the app's own page, as a browser sends it: the browser marks it
    # same-origin even when it reports the Origin as "null".
    own_form = client.post(
        "/screening",
        data={"sni": "6210"},
        headers={"origin": "null", "sec-fetch-site": "same-origin"},
    )
    assert own_form.status_code == 200 and "fem siffror" in own_form.text
    same_site = client.post("/api/test/scb", headers={"origin": "http://testserver"})
    assert same_site.status_code == 200 and same_site.json()["ok"]
    assert client.post("/api/test/bolagsverket").json()["ok"]
    assert client.get("/analyser/..%2F..%2Fetc/sammanfattning").status_code == 404

    # Not framable, scripts and styles only from the app, nothing cached or referred.
    headers = client.get("/").headers
    assert headers["x-frame-options"] == "DENY"
    assert "frame-ancestors 'none'" in headers["content-security-policy"]
    assert "script-src 'self'" in headers["content-security-policy"]
    assert headers["x-content-type-options"] == "nosniff"
    assert headers["referrer-policy"] == "same-origin"  # nothing is sent to other sites
    assert "<script>" not in client.get("/").text  # no inline script for the policy to block

    # An oversized upload is refused before its body is read.
    too_large = client.post(
        "/analyser", content=b"x", headers={"content-length": str(60 * 1024 * 1024)}
    )
    assert too_large.status_code == 413


def test_unknown_addresses_and_bad_input_read_well(client: TestClient) -> None:
    unknown = client.get("/finns-inte")
    assert unknown.status_code == 404 and "Sidan finns inte" in unknown.text
    assert "<html" in unknown.text  # the app's own page, not a JSON body
    bad_number = client.post("/api/analyser/20990101-000000-abcdef/urval", content=b"{x")
    assert bad_number.status_code == 400 and "error" in bad_number.json()
    gone = client.get("/api/jobb/0000")
    assert gone.status_code == 404 and "Starta den igen" in gone.json()["error"]

    # Odd form values are dropped or refused readably, and the choices are kept.
    odd = client.post(
        "/screening",
        data={"sni": "62100, 62100", "oms_kl": ["x", "6"], "lan": ["1", "evil"], "topp": "abc"},
    )
    assert odd.status_code == 200 and "Screening" in odd.text
    form_error = client.post("/screening", data={"sni": "6210", "oms_kl": ["6"], "lan": ["1"]})
    assert "fem siffror" in form_error.text
    assert 'name="oms_kl" value="6" checked' in form_error.text
    assert 'name="lan" value="01" checked' in form_error.text


def test_memo_html_keeps_negative_amounts() -> None:
    html = views.memo_html("- -600,0 TSEK periodisering\n-3,0 % lägre marginal\n## Rubrik")
    assert "<li>-600,0 TSEK periodisering</li>" in html
    assert "<p>-3,0 % lägre marginal</p>" in html and "<h3>Rubrik</h3>" in html


def test_demo_company_runs_from_the_form(client: TestClient) -> None:
    page = client.post("/analyser", data={"demo": "1"})
    assert page.url.path.endswith("/sammanfattning") and "Demobolaget" in page.text
    assert "Utan AI-klassning" in page.text  # no API key in the test settings


class _Throttled:
    """Bolagsverket stand-in that answers 429 until `state["calm"]` is set."""

    def __init__(self, state: dict[str, bool]) -> None:
        self.state = state

    def __enter__(self) -> "_Throttled":
        return self

    def __exit__(self, *exc: object) -> None:
        return None

    def latest_report(self, org_nr: str) -> None:
        if not self.state["calm"]:
            raise BolagsverketUnavailable("Bolagsverket begränsar antalet anrop just nu.")
        return None


def test_screening_can_fetch_what_bolagsverket_did_not_answer(tmp_path: Path, demo: Path) -> None:
    units, turnover = _units()
    state = {"calm": False}
    app = create_app(
        settings_provider=_settings,
        workspace=tmp_path / "arbetsyta",
        scb_factory=lambda s: fake_scb(units, turnover),  # type: ignore[arg-type]
        bolagsverket_factory=lambda s: _Throttled(state),  # type: ignore[arg-type,return-value]
        inline_jobs=True,
        demo_file=demo,
        allowed_hosts=("testserver",),
    )
    with TestClient(app) as client:
        throttled = client.post(
            "/screening", data={"sni": "62100", "oms_kl": ["6"], "med_bokslut": "1"}
        )
        assert "Hämta återstående" in throttled.text and "Kunde inte hämtas" in throttled.text
        item_id = throttled.url.path.split("/")[2]

        state["calm"] = True
        completed = client.post(f"/screening/{item_id}/komplettera")
        assert completed.url.path == f"/screening/{item_id}"
        assert "Hämta återstående" not in completed.text
        assert "Kunde inte hämtas" not in completed.text and "Saknas digitalt" in completed.text


def test_running_and_failed_jobs_stay_visible(tmp_path: Path, demo: Path) -> None:
    """Leaving the progress page must not hide a run: overview, lists and the menu show it."""
    release = threading.Event()
    app = create_app(
        settings_provider=_settings,
        workspace=tmp_path / "arbetsyta",
        demo_file=demo,
        allowed_hosts=("testserver",),
    )
    jobs = app.state.ctx.jobs

    def slow(report: Reporter) -> str:
        report.stage("Hämtar bolag i vald bransch")
        release.wait(10)
        return "/screening"

    def broken(report: Reporter) -> str:
        raise ValueError("SCB svarar inte just nu.")

    with TestClient(app) as client:
        running = jobs.start("screening", "Screening 62100", ["Hämtar bolag i vald bransch"], slow)
        failed = jobs.start("analys", "Analyserar fel.se", ["Läser SIE-filen"], broken)
        for _ in range(100):
            if jobs.get(failed.id).status == "fel":
                break
            time.sleep(0.02)

        overview = client.get("/").text
        assert "Screening 62100" in overview and f"/jobb/{running.id}" in overview
        assert "Hämtar bolag i vald bransch" in overview  # the stage it is in
        assert "Analyserar fel.se" in overview and "SCB svarar inte just nu." in overview
        assert "1 körning pågår" in client.get("/metod").text  # the menu, on any page
        assert "Screening 62100" in client.get("/screening").text
        assert "Screening 62100" not in client.get("/analyser").text
        assert "Analyserar fel.se" in client.get("/analyser").text

        release.set()
        for _ in range(100):
            if jobs.get(running.id).status == "klar":
                break
            time.sleep(0.02)
        assert "1 körning pågår" not in client.get("/metod").text


def test_app_command_listens_where_told(monkeypatch: pytest.MonkeyPatch) -> None:
    """Locally on 127.0.0.1 only; in a container on every interface, behind Docker's port."""
    import socket

    import uvicorn
    from typer.testing import CliRunner

    from diligens.cli import app as cli_app

    calls: list[dict[str, object]] = []
    monkeypatch.setattr(uvicorn, "run", lambda *a, **kw: calls.append(kw))
    with socket.socket() as spare:
        spare.bind(("127.0.0.1", 0))
        port = spare.getsockname()[1]

    local = CliRunner().invoke(cli_app, ["app", "--port", str(port), "--inte-oppna"])
    assert local.exit_code == 0, local.output
    container = CliRunner().invoke(
        cli_app, ["app", "--port", str(port), "--host", "0.0.0.0", "--inte-oppna"]
    )
    assert container.exit_code == 0, container.output
    assert [c["host"] for c in calls] == ["127.0.0.1", "0.0.0.0"]

    with socket.socket() as busy:
        busy.bind(("127.0.0.1", port))
        busy.listen()
        taken = CliRunner().invoke(cli_app, ["app", "--port", str(port), "--inte-oppna"])
    assert taken.exit_code == 1 and "används redan" in taken.output
