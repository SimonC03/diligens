"""M6 screening tests: filters and fit score (pure) and the full run with fake APIs."""

import datetime as dt
import json
from decimal import Decimal
from pathlib import Path
from typing import Any

import httpx
import pytest
from openpyxl import load_workbook
from test_scb import unit as scb_unit
from typer.testing import CliRunner

import diligens.cli as cli
from diligens.analysis.screening import financial_score, profile_score, select_candidates
from diligens.cli import app
from diligens.models import AnnualReport, Figure, ScbUnit, ScreeningCriteria
from diligens.sources.bolagsverket import BolagsverketClient
from diligens.sources.scb import ScbClient, ScbError


def _unit(
    n: int,
    *,
    anst: str = "5",
    lan: str = "01",
    jur: str = "49",
    stat: str = "1",
    sni: str = "62100",
) -> ScbUnit:
    return ScbUnit(
        peOrgNr=f"1655990{n:05d}",
        orgNr=f"55990{n:05d}",
        namn=f"Testbolag {n} AB",
        lanSate=lan,
        kommunSate="0180",
        anstKl=anst,
        ftgStat=stat,
        jurform=jur,
        primary_sni=sni,
    )


def _report(revenue: str, prior: str | None, ebit: str) -> AnnualReport:
    figures = {
        "nettoomsattning": Figure(value=Decimal(revenue), source_ref="iXBRL Nettoomsattning"),
        "rorelseresultat": Figure(value=Decimal(ebit), source_ref="iXBRL Rorelseresultat"),
    }
    if prior is not None:
        figures["nettoomsattning_fg"] = Figure(value=Decimal(prior), source_ref="iXBRL fg")
    return AnnualReport(
        document_id="doc",
        org_nr="5599000001",
        period_start=dt.date(2025, 1, 1),
        period_end=dt.date(2025, 12, 31),
        figures=figures,
    )


CRITERIA = ScreeningCriteria(
    sni=["62100"], turnover_classes=["6"], employee_classes=["5"], counties=["01"]
)


def test_filters_and_profile_score() -> None:
    units = [
        _unit(1),  # perfect match
        _unit(2, anst="4", lan="12"),  # adjacent employee class, other county
        _unit(3, jur="10"),  # sole trader: not an aktiebolag
        _unit(4, stat="9"),  # no longer active
        _unit(5, anst="8"),  # far outside the employee class
        _unit(6),  # turnover class adjacent (7)
        _unit(7),  # turnover class far outside (9), dropped
    ]
    turnover = {
        "165599000001": "6",
        "165599000002": "6",
        "165599000003": "6",
        "165599000004": "6",
        "165599000005": "6",
        "165599000006": "7",
        "165599000007": "9",
    }

    candidates = select_candidates(units, turnover, CRITERIA)

    assert [c.org_nr for c in candidates] == ["5599000001", "5599000006", "5599000002"]  # by score
    best = {c.org_nr: c for c in candidates}
    assert best["5599000001"].profile_points == 60
    assert best["5599000002"].profile_points == 15 + 15 + 5 + 0 + 10
    assert best["5599000006"].profile_points == 15 + 7 + 10 + 10 + 10
    reasons = " ".join(best["5599000002"].profile_reasons)
    assert "anställdaklass" in reasons.lower() and "län" in reasons.lower()

    points, reasons = profile_score(_unit(9, lan="03"), "6", ScreeningCriteria(sni=["62100"]))
    assert points == 60  # no size or county targets: nothing to miss
    assert reasons


def test_financial_score_and_missing_reports() -> None:
    strong = financial_score(_report("33000000", "29000000", "3600000"))
    assert strong.points == 40
    assert (
        strong.growth is not None
        and strong.growth.value == Decimal("33000000") / Decimal("29000000") - 1
    )
    assert strong.ebit_margin is not None and "Rorelseresultat" in strong.ebit_margin.source_ref

    middle = financial_score(_report("30000000", "29000000", "2000000"))  # +3 %, 6,7 %
    assert middle.points == 10 + 10
    weak = financial_score(_report("27000000", "29000000", "500000"))  # shrinking, 1,9 %
    assert weak.points == 0
    no_prior = financial_score(_report("27000000", None, "3000000"))
    assert no_prior.points == 20 and no_prior.growth is None
    assert any("föregående år" in r for r in no_prior.reasons)

    missing = financial_score(None)
    assert missing.points is None and missing.status == "Saknas digitalt"


FIXTURES = Path(__file__).parent / "fixtures"


def _fake_scb(units: list[dict[str, Any]], turnover: dict[str, list[str]]) -> ScbClient:
    tables = {
        name: json.loads((FIXTURES / "scb" / f"{name}.json").read_text(encoding="utf-8"))
        for name in ("omsklkoder", "anstklkoder", "lankoder")
    }

    def handle(request: httpx.Request) -> httpx.Response:
        path = request.url.path
        if path.startswith("/v1/kodtabeller/"):
            return httpx.Response(200, json=tables[path.rsplit("/", 1)[-1]])
        if path == "/v1/juridiskaenheter/naringsgren/62100":
            jes = units
        elif path.startswith("/v1/juridiskaenheter/omsattningsklass/"):
            wanted = turnover.get(path.rsplit("/", 1)[-1], [])
            jes = [u for u in units if u["peOrgNr"] in wanted]
        else:
            return httpx.Response(404, json={"status": 404, "detail": "x"})
        return httpx.Response(
            200,
            json={"jes": jes, "pagination": {"nextCursorId": 0, "limit": 5000, "hasMore": False}},
        )

    return ScbClient(
        "k",
        transport=httpx.MockTransport(handle),
        sleep=lambda s: None,
        now=lambda: dt.datetime(2026, 10, 17, 12, 0),
    )


def _fake_registry() -> BolagsverketClient:
    package = (FIXTURES / "bolagsverket" / "dokument_exempel.zip").read_bytes()

    def handle(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/oauth2/token":
            return httpx.Response(200, json={"access_token": "t", "expires_in": 3600})
        if request.url.path.endswith("/dokumentlista"):
            org = json.loads(request.content)["identitetsbeteckning"]
            docs = (
                [
                    {
                        "dokumentId": "d1_paket",
                        "rapporteringsperiodTom": "2025-12-31",
                        "registreringstidpunkt": "2026-06-30",
                    }
                ]
                if org == "5599000001"
                else []
            )
            return httpx.Response(200, json={"dokument": docs})
        return httpx.Response(200, content=package)

    return BolagsverketClient("i", "s", transport=httpx.MockTransport(handle))


def test_screen_writes_long_list_with_sources_and_missing_reports(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    units = [
        scb_unit(1),
        scb_unit(2, lan="12"),
        scb_unit(3, jur="10"),  # sole trader, filtered out
        scb_unit(4, stat="9"),  # inactive, filtered out
        scb_unit(5, anst="8"),  # too large, filtered out
        scb_unit(6),
    ]
    pe = {n: f"1655990{n:05d}" for n in range(1, 7)}
    turnover = {"6": [pe[1], pe[2], pe[3], pe[4], pe[5]], "7": [pe[6]]}
    monkeypatch.setattr(cli, "make_scb_client", lambda s: _fake_scb(units, turnover))
    monkeypatch.setattr(cli, "make_bolagsverket_client", lambda s: _fake_registry())

    run = CliRunner().invoke(
        app,
        [
            "screen",
            "--sni",
            "62100",
            "--oms-kl",
            "6",
            "--anst-kl",
            "5",
            "--lan",
            "1",
            "--out",
            str(tmp_path),
        ],
    )

    assert run.exit_code == 0, run.output
    assert "3 kandidater av 6 bolag" in run.output
    wb = load_workbook(tmp_path / "long_list.xlsx")
    assert wb.sheetnames == ["Long list", "Kriterier"]
    ws = wb["Long list"]
    header = [c.value for c in ws[1]]
    rows = [
        dict(zip(header, (c.value for c in row), strict=True)) for row in ws.iter_rows(min_row=2)
    ]
    assert [r["Org.nr"] for r in rows] == ["559900-0001", "559900-0006", "559900-0002"]

    best = rows[0]
    assert best["Profil"] == 60 and best["Nyckeltal"] == 40  # growth 12,3 %, margin 10,0 %
    assert best["Poäng"] == "=C2+D2"
    assert best["Bokslut"] == "Årsredovisning 2025"
    assert best["Nettoomsättning TSEK"] == pytest.approx(32450.0)
    assert best["Tillväxt"] == pytest.approx(32450 / 28900 - 1)
    assert "iXBRL se-gen-base:Nettoomsattning" in best["Källa"] and "d1_paket" in best["Källa"]
    assert "SCB AFR" in best["Källa"]
    assert best["Omsättningsklass"] == "20 000 - 49 999 tkr" and best["Län"] == "Stockholm"

    # No digital annual report: shown as such, no figures guessed, profile score only.
    for row in rows[1:]:
        assert row["Bokslut"] == "Saknas digitalt"
        assert row["Nettoomsättning TSEK"] is None and row["Nyckeltal"] == 0
    assert "en klass från målet" in rows[1]["Skäl"]  # turnover class 7 next to 6
    assert "annat län" in rows[2]["Skäl"]

    criteria = {r[0].value: r[1].value for r in wb["Kriterier"].iter_rows(min_row=2)}
    assert criteria["Omsättningsklass"] == "6 20 000 - 49 999 tkr"
    assert criteria["Län"] == "01 Stockholm"

    # Missing keys give a readable error.
    def no_key(settings: Any) -> ScbClient:
        raise ScbError("SCB_API_KEY saknas i .env.")

    monkeypatch.setattr(cli, "make_scb_client", no_key)
    failed = CliRunner().invoke(app, ["screen", "--sni", "62100", "--out", str(tmp_path)])
    assert failed.exit_code == 1 and "SCB_API_KEY" in failed.output
